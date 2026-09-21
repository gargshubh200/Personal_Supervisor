import os
import json
from pathlib import Path
from typing import Dict, Any, List, Optional
from datetime import datetime
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from mcp.server.mcpserver import MCPServer
from langfuse import observe, get_client

# Google API Client Imports
from google.oauth2 import service_account
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload
import google.auth

mcp = MCPServer("DispatchServer")
load_dotenv()

OUTPUT_DIR = Path(__file__).parent / "gcs_storage" / "output_resumes"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive"]


def _load_profile() -> Dict[str, Any]:
    """Locates and loads master_profile.json from project paths."""
    candidate_paths = [
        Path(__file__).parent.parent / "master_profile.json",
        Path(__file__).parent / "master_profile.json",
        Path("master_profile.json")
    ]
    for p in candidate_paths:
        if p.exists():
            with open(p, "r", encoding="utf-8") as f:
                return json.load(f)

    try:
        from OtherMCP.ground_truth_mcp import _load_profile as load_gt
        return load_gt()
    except Exception as e:
        raise FileNotFoundError(f"master_profile.json not found in candidate paths. Error: {str(e)}")


def set_spacing(paragraph, space_before=0, space_after=2, line_spacing=1.15):
    p_format = paragraph.paragraph_format
    p_format.space_before = Pt(space_before)
    p_format.space_after = Pt(space_after)
    p_format.line_spacing = line_spacing


# ------------------------------------------------------------------
# Google Drive Authentication & Upload Service
# ------------------------------------------------------------------

def _get_drive_service():
    creds = None
    service_account_path = os.getenv("GOOGLE_APPLICATION_CREDENTIALS", "service_account.json")
    token_path = Path(__file__).parent / "token.json"
    credentials_path = Path(__file__).parent / "credentials.json"

    # 1. Try local service account file if present
    if os.path.exists(service_account_path):
        try:
            creds = service_account.Credentials.from_service_account_file(
                service_account_path, scopes=DRIVE_SCOPES
            )
        except Exception as e:
            print(f"⚠️ Service account file auth failed: {e}")

    # 2. Try local user tokens (OAuth)
    elif token_path.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(token_path), DRIVE_SCOPES)
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
                with open(token_path, "w", encoding="utf-8") as token_file:
                    token_file.write(creds.to_json())
        except Exception as e:
            print(f"⚠️ Token auth failed: {e}")
            creds = None

    elif credentials_path.exists():
        try:
            flow = InstalledAppFlow.from_client_secrets_file(str(credentials_path), DRIVE_SCOPES)
            creds = flow.run_local_server(port=0)
            with open(token_path, "w", encoding="utf-8") as token_file:
                token_file.write(creds.to_json())
        except Exception as e:
            print(f"⚠️ Credentials auth failed: {e}")
            creds = None

    # 3. Fallback to Cloud Run / GCP Application Default Credentials (ADC)
    if not creds:
        try:
            creds, _ = google.auth.default(scopes=DRIVE_SCOPES)
            print("🔐 Authenticated with GCP Application Default Credentials.")
        except Exception as e:
            print(f"⚠️ ADC authentication failed: {str(e)}")
            return None

    return build("drive", "v3", credentials=creds)


def _upload_to_google_drive(file_path: Path, convert_to_gdoc: bool = True) -> Dict[str, str]:
    try:
        service = _get_drive_service()
        if not service:
            print("⚠️ Google Drive service not initialized. Skipping Drive upload.")
            return {"status": "skipped", "drive_url": None}

        folder_id = os.getenv("GOOGLE_DRIVE_FOLDER_ID")
        file_metadata: Dict[str, Any] = {"name": file_path.name}

        if folder_id:
            file_metadata["parents"] = [folder_id]

        if convert_to_gdoc:
            file_metadata["mimeType"] = "application/vnd.google-apps.document"

        media = MediaFileUpload(
            str(file_path),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            resumable=True
        )

        drive_file = service.files().create(
            body=file_metadata,
            media_body=media,
            fields="id, webViewLink"
        ).execute()

        file_id = drive_file.get("id")
        web_link = drive_file.get("webViewLink")

        try:
            service.permissions().create(
                fileId=file_id,
                body={'type': 'anyone', 'role': 'reader'}
            ).execute()
        except Exception as perm_err:
            print(f"⚠️ Permission grant warning: {str(perm_err)}")

        print(f"☁️ Google Drive Upload Successful! File ID: {file_id}")
        return {
            "status": "success",
            "file_id": file_id,
            "drive_url": web_link
        }

    except Exception as e:
        print(f"❌ Google Drive Upload Error: {str(e)}")
        return {"status": "error", "error": str(e), "drive_url": None}


# ------------------------------------------------------------------
# Resume Generation Tool
# ------------------------------------------------------------------

@observe(name="Dispatch: Generate ATS Resume Document")
@mcp.tool()
def generate_tailored_resume_docx(
        company_name: str,
        role_title: str,
        tailored_summary: str,
        tailored_bullets: List[Dict[str, Any]]
) -> Dict[str, str]:
    profile = _load_profile()
    info = profile["personal_info"]

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    name_p = doc.add_paragraph()
    set_spacing(name_p, space_after=2)
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = name_p.add_run(info["name"].upper())
    run_name.font.name = "Arial"
    run_name.font.size = Pt(16)
    run_name.font.bold = True

    contact_p = doc.add_paragraph()
    set_spacing(contact_p, space_after=10)
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_text = f"+91 9109268883 | {info['email']} | {info['linkedin']} | {info['github']}"
    run_contact = contact_p.add_run(contact_text)
    run_contact.font.name = "Arial"
    run_contact.font.size = Pt(9.5)

    def add_section_heading(title: str):
        p = doc.add_paragraph()
        set_spacing(p, space_before=8, space_after=4)
        run = p.add_run(title.upper())
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        return p

    add_section_heading("Summary")
    sum_p = doc.add_paragraph()
    set_spacing(sum_p, space_after=8)
    sum_run = sum_p.add_run(tailored_summary)
    sum_run.font.name = "Arial"
    sum_run.font.size = Pt(10)

    add_section_heading("Experience")
    bullet_map = {b.get("original_bullet", ""): b.get("tailored_bullet", "") for b in tailored_bullets}

    for exp in profile.get("experience", []):
        job_header = doc.add_paragraph()
        set_spacing(job_header, space_before=4, space_after=2)

        role_run = job_header.add_run(f"{exp['role']} — {exp['company']}")
        role_run.font.name = "Arial"
        role_run.font.size = Pt(10.5)
        role_run.font.bold = True

        details_text = f" | {exp['period']} | {exp.get('location', '')}"
        details_run = job_header.add_run(details_text)
        details_run.font.name = "Arial"
        details_run.font.size = Pt(10)

        # Experience items may either group bullets under named "projects" or
        # list flat "achievements" directly (both fields are Optional in the
        # master profile schema, e.g. for roles without distinct projects).
        for project in (exp.get("projects") or []):
            proj_p = doc.add_paragraph()
            set_spacing(proj_p, space_before=2, space_after=2)
            proj_run = proj_p.add_run(project["name"])
            proj_run.font.name = "Arial"
            proj_run.font.size = Pt(10)
            proj_run.font.bold = True
            proj_run.font.italic = True

            for orig_bullet in (project.get("achievements") or []):
                bullet_text = bullet_map.get(orig_bullet, orig_bullet)
                bp = doc.add_paragraph(style='List Bullet')
                set_spacing(bp, space_before=0, space_after=2)
                b_run = bp.add_run(bullet_text)
                b_run.font.name = "Arial"
                b_run.font.size = Pt(9.5)

        for orig_bullet in (exp.get("achievements") or []):
            bullet_text = bullet_map.get(orig_bullet, orig_bullet)
            bp = doc.add_paragraph(style='List Bullet')
            set_spacing(bp, space_before=0, space_after=2)
            b_run = bp.add_run(bullet_text)
            b_run.font.name = "Arial"
            b_run.font.size = Pt(9.5)

    add_section_heading("Skills")
    skills = profile.get("skills", {})

    for category, skill_list in skills.items():
        sk_p = doc.add_paragraph()
        set_spacing(sk_p, space_before=0, space_after=2)
        cat_formatted = category.replace("_", " ").title()

        cat_run = sk_p.add_run(f"{cat_formatted}: ")
        cat_run.font.name = "Arial"
        cat_run.font.size = Pt(9.5)
        cat_run.font.bold = True

        list_run = sk_p.add_run(", ".join(skill_list))
        list_run.font.name = "Arial"
        list_run.font.size = Pt(9.5)

    standalone_projects = profile.get("projects") or []
    if standalone_projects:
        add_section_heading("Projects")
        for proj in standalone_projects:
            proj_header = doc.add_paragraph()
            set_spacing(proj_header, space_before=4, space_after=1)

            header_text = proj.get("name", "")
            tech_stack = proj.get("tech_stack") or []
            if tech_stack:
                header_text += f" ({', '.join(tech_stack)})"

            proj_name_run = proj_header.add_run(header_text)
            proj_name_run.font.name = "Arial"
            proj_name_run.font.size = Pt(10.5)
            proj_name_run.font.bold = True

            description = proj.get("description")
            if description:
                desc_p = doc.add_paragraph()
                set_spacing(desc_p, space_before=0, space_after=2)
                desc_run = desc_p.add_run(description)
                desc_run.font.name = "Arial"
                desc_run.font.size = Pt(9.5)
                desc_run.font.italic = True

            for ach in (proj.get("achievements") or []):
                ach_p = doc.add_paragraph(style='List Bullet')
                set_spacing(ach_p, space_before=0, space_after=2)
                ach_run = ach_p.add_run(ach)
                ach_run.font.name = "Arial"
                ach_run.font.size = Pt(9.5)

    add_section_heading("Education")
    edu_data = profile.get("education", {})
    edu_p = doc.add_paragraph()
    set_spacing(edu_p, space_before=2, space_after=2)

    if isinstance(edu_data, dict):
        degree = edu_data.get("degree", "")
        inst = edu_data.get("institution", "")
        period = edu_data.get("period", "")
        edu_text = f"{degree} — {inst} ({period})" if period else f"{degree} — {inst}"

        edu_run = edu_p.add_run(edu_text)
        edu_run.font.name = "Arial"
        edu_run.font.size = Pt(9.5)

        for ach in edu_data.get("achievements", []):
            ach_p = doc.add_paragraph(style='List Bullet')
            set_spacing(ach_p, space_before=0, space_after=2)
            ach_run = ach_p.add_run(ach)
            ach_run.font.name = "Arial"
            ach_run.font.size = Pt(9)
    else:
        edu_run = edu_p.add_run(str(edu_data))
        edu_run.font.name = "Arial"
        edu_run.font.size = Pt(9.5)

    safe_company = "".join([c for c in company_name if c.isalnum() or c in (" ", "_")]).strip()
    safe_role = "".join([c for c in role_title if c.isalnum() or c in (" ", "_")]).strip()
    file_name = f"ATS_Resume_Sahil_Garg_{safe_company}_{safe_role}.docx".replace(" ", "_")
    output_path = OUTPUT_DIR / file_name

    doc.save(str(output_path))
    print(f"📄 Local Resume Saved: {output_path}")

    drive_result = _upload_to_google_drive(output_path, convert_to_gdoc=True)

    try:
        langfuse = get_client()
        langfuse.update_current_span(
            metadata={
                "company_name": company_name,
                "role_title": role_title,
                "local_path": str(output_path),
                "drive_status": drive_result["status"],
                "drive_url": str(drive_result.get("drive_url"))
            }
        )
    except Exception:
        pass

    return {
        "local_path": str(output_path),
        "drive_url": drive_result.get("drive_url") or "Upload skipped or failed",
        "file_id": drive_result.get("file_id", "")
    }


# ------------------------------------------------------------------
# Cover Letter Generation Tool
# ------------------------------------------------------------------

@observe(name="Dispatch: Generate Cover Letter Document")
@mcp.tool()
def generate_cover_letter_docx(
        company_name: str,
        role_title: str,
        cover_letter_data: Dict[str, Any]
) -> Dict[str, str]:
    """
    Renders structured cover letter output into an executive DOCX document,
    saves it locally, and uploads it to Google Drive.
    """
    profile = _load_profile()
    info = profile["personal_info"]

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Header Block
    name_p = doc.add_paragraph()
    set_spacing(name_p, space_after=2)
    run_name = name_p.add_run(info["name"].upper())
    run_name.font.name = "Arial"
    run_name.font.size = Pt(14)
    run_name.font.bold = True

    contact_p = doc.add_paragraph()
    set_spacing(contact_p, space_after=12)
    contact_text = f"{info['email']} | +91 9109268883 | {info['linkedin']} | {info['github']}"
    run_contact = contact_p.add_run(contact_text)
    run_contact.font.name = "Arial"
    run_contact.font.size = Pt(9.5)

    # 2. Date Block
    date_p = doc.add_paragraph()
    set_spacing(date_p, space_after=12)
    date_run = date_p.add_run(datetime.now().strftime("%B %d, %Y"))
    date_run.font.name = "Arial"
    date_run.font.size = Pt(10)

    # 3. Salutation
    salutation_p = doc.add_paragraph()
    set_spacing(salutation_p, space_after=10)
    sal_run = salutation_p.add_run(cover_letter_data.get("salutation", f"Dear Hiring Team at {company_name},"))
    sal_run.font.name = "Arial"
    sal_run.font.size = Pt(10.5)
    sal_run.font.bold = True

    # 4. Opening Paragraph
    opening_p = doc.add_paragraph()
    set_spacing(opening_p, space_after=10)
    op_run = opening_p.add_run(cover_letter_data.get("opening_paragraph", ""))
    op_run.font.name = "Arial"
    op_run.font.size = Pt(10)

    # 5. Body Paragraphs
    for body_para in cover_letter_data.get("body_paragraphs", []):
        bp = doc.add_paragraph()
        set_spacing(bp, space_after=10)
        bp_run = bp.add_run(body_para)
        bp_run.font.name = "Arial"
        bp_run.font.size = Pt(10)

    # 6. Closing Paragraph
    closing_p = doc.add_paragraph()
    set_spacing(closing_p, space_after=16)
    cp_run = closing_p.add_run(cover_letter_data.get("closing_paragraph", ""))
    cp_run.font.name = "Arial"
    cp_run.font.size = Pt(10)

    # 7. Sign-off
    sign_p = doc.add_paragraph()
    set_spacing(sign_p, space_after=2)
    s_run1 = sign_p.add_run("Sincerely,\n\n")
    s_run1.font.name = "Arial"
    s_run1.font.size = Pt(10)

    s_run2 = sign_p.add_run(info["name"])
    s_run2.font.name = "Arial"
    s_run2.font.size = Pt(10.5)
    s_run2.font.bold = True

    # Save to local file
    safe_company = "".join([c for c in company_name if c.isalnum() or c in (" ", "_")]).strip()
    safe_role = "".join([c for c in role_title if c.isalnum() or c in (" ", "_")]).strip()
    file_name = f"Cover_Letter_Sahil_Garg_{safe_company}_{safe_role}.docx".replace(" ", "_")
    output_path = OUTPUT_DIR / file_name

    doc.save(str(output_path))
    print(f"📄 Local Cover Letter Saved: {output_path}")

    # Upload to Google Drive
    drive_result = _upload_to_google_drive(output_path, convert_to_gdoc=True)

    try:
        langfuse = get_client()
        langfuse.update_current_span(
            metadata={
                "company_name": company_name,
                "role_title": role_title,
                "local_path": str(output_path),
                "drive_status": drive_result["status"],
                "drive_url": str(drive_result.get("drive_url"))
            }
        )
    except Exception:
        pass

    return {
        "local_path": str(output_path),
        "drive_url": drive_result.get("drive_url") or "Upload skipped or failed",
        "file_id": drive_result.get("file_id", "")
    }